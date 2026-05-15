from pathlib import Path

from pptx import Presentation

from .base import DocumentParser


class PptxParser(DocumentParser):
    def extract_text(self, file_path: Path) -> str:
        prs = Presentation(str(file_path))
        texts: list[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        texts.append(paragraph.text)
        return "\n".join(texts)
